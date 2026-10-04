"""Read supplied sources without changing them; preserve page/paragraph locators."""
import hashlib
import json
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "r1" / "source-extracts"
OUT.mkdir(parents=True, exist_ok=True)
KEYWORDS = ("经验", "升级", "评分", "评价", "速杀", "无伤", "成长")
manifest = []
for path in sorted((ROOT / "docs").iterdir()):
    if path.suffix.lower() not in (".pdf", ".docx"):
        continue
    record = {"file": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "sections": []}
    if path.suffix.lower() == ".pdf":
        doc = PdfReader(path)
        for i, page in enumerate(doc.pages, 1):
            record["sections"].append({"locator": f"PDF page {i}", "text": page.extract_text() or ""})
    else:
        doc = Document(path)
        # Paragraph numbers refer to document.paragraphs (including empty paragraphs).
        for i, p in enumerate(doc.paragraphs, 1):
            if p.text.strip():
                record["sections"].append({"locator": f"paragraph {i}", "text": p.text})
        for ti, table in enumerate(doc.tables, 1):
            for ri, row in enumerate(table.rows, 1):
                record["sections"].append({"locator": f"table {ti}, row {ri}", "text": " | ".join(c.text for c in row.cells)})
    (OUT / (path.stem + ".json")).write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
    hits = [section for section in record["sections"] if any(word in section["text"] for word in KEYWORDS)]
    manifest.append({"file": record["file"], "sha256": record["sha256"], "sections": len(record["sections"]), "hits": len(hits)})
    (OUT / (path.stem + "-hits.txt")).write_text("\n\n".join(section["locator"] + "\n" + section["text"] for section in hits), encoding="utf-8")
(OUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(manifest, ensure_ascii=False, indent=2))
